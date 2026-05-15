#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
collector/registry_manager.py
Управление реестром должников: совместимый shim поверх CRM.

Версия: 1.1.0 (2026-05-15)
"""

import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

from dotenv import load_dotenv
from zoneinfo import ZoneInfo

from bot.crm_clients import (
    load_clients as crm_load_clients,
    save_clients as crm_save_clients,
    set_client_alias,
    set_client_phone,
    load_contacts_for_collector,
)
from collector.logging_utils import get_collector_logger

load_dotenv(
    dotenv_path=Path(__file__).resolve().parent.parent / ".env",
    encoding="utf-8-sig",
    override=False,
)

TZ = ZoneInfo(os.getenv("TZ", "Asia/Almaty"))

ROOT_DIR = Path(__file__).resolve().parent.parent
CONTACTS_PATH = ROOT_DIR / "config" / "debtors_contacts.json"
REGISTRY_XLSX = ROOT_DIR / "config" / "debtors_registry.xlsx"

logger = get_collector_logger(__name__)


def _load_contacts() -> Dict[str, Any]:
    """Compatibility view: CRM is primary, legacy file is read-only fallback."""
    try:
        return load_contacts_for_collector()
    except Exception as e:
        logger.error("Ошибка чтения unified collector contacts: %s", e)
        return {}


def auto_register_client(
    name: str,
    manager: str = "",
    amount: float = 0.0,
    days: int = 0,
    violation: bool = False,
) -> bool:
    """
    Автоматически добавляет клиента в CRM без контактных данных.

    Телефон и Telegram ID остаются пустыми — менеджер должен заполнить через бота.
    Возвращает True если клиент был добавлен (новый), False если уже существует.
    """
    data = crm_load_clients()
    contacts = data.get("clients", {})

    if name in contacts:
        return False
    name_lower = name.lower().strip()
    for key in contacts:
        if key.lower().strip() == name_lower:
            return False

    contacts[name] = {
        "whatsapp": "",
        "telegram_id": "",
        "manager": manager,
        "language": "ru",
        "do_not_call": False,
        "_auto_registered": True,
        "_needs_phone": True,
        "_debt_amount": round(amount, 2),
        "_days_overdue": days,
        "_violation": violation,
        "_registered_date": datetime.now(tz=TZ).date().isoformat(),
    }

    data["clients"] = contacts
    if not crm_save_clients(data):
        return False

    logger.info(
        "Авторегистрация: %s (менеджер=%s, долг=%.0f, дней=%d)",
        name,
        manager,
        amount,
        days,
    )
    export_registry_excel(contacts)
    return True


def update_client_phone(name: str, phone: str) -> bool:
    """Обновляет номер WhatsApp клиента через CRM."""
    if not set_client_phone(name, phone, reviewer="registry_manager"):
        return False

    logger.info(
        "Телефон обновлён для клиента: ***%s",
        phone[-4:] if len(phone) >= 4 else "****",
    )
    export_registry_excel()
    return True


def update_client_language(name: str, language: str) -> bool:
    """Устанавливает язык общения с клиентом ('ru' или 'kz') через CRM."""
    if language not in ("ru", "kz"):
        logger.warning("update_client_language: неверный язык: %s", language)
        return False

    data = crm_load_clients()
    contacts = data.get("clients", {})
    if name not in contacts:
        logger.warning("update_client_language: клиент не найден: %s", name)
        return False
    contacts[name]["language"] = language
    data["clients"] = contacts
    if not crm_save_clients(data):
        return False
    logger.info("Язык обновлён для %s: %s", name, language)
    export_registry_excel()
    return True


def update_client_display_name(name: str, display_name: str) -> bool:
    """Сохраняет откорректированное отображаемое имя клиента через CRM."""
    if not set_client_alias(name, display_name.strip()):
        return False

    logger.info("display_name обновлён для %s: %s", name, display_name.strip())
    export_registry_excel()
    return True


def export_registry_excel(contacts: Optional[Dict[str, Any]] = None) -> bool:
    """
    Экспортирует реестр должников в Excel для просмотра администратором.

    Файл: config/debtors_registry.xlsx
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl не установлен — Excel экспорт недоступен")
        return False

    if contacts is None:
        contacts = _load_contacts()

    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр должников"

    headers = [
        "Статус",
        "Клиент (1С)",
        "Имя для сообщений",
        "WhatsApp",
        "Telegram ID",
        "Менеджер",
        "Язык",
        "Не звонить",
        "Долг (тг)",
        "Дней просрочки",
        "Нарушение",
        "Дата добавления",
    ]
    header_fill = PatternFill("solid", fgColor="1A3A5C")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    green_fill = PatternFill("solid", fgColor="E2EFDA")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    bold_font = Font(bold=True)

    for row_idx, (name, info) in enumerate(contacts.items(), 2):
        if not isinstance(info, dict):
            continue
        needs_phone = info.get("_needs_phone", False) or not info.get("whatsapp")
        status = "⚠️ Нет телефона" if needs_phone else "✅ Готов"
        row_fill = orange_fill if needs_phone else green_fill

        values = [
            status,
            name,
            info.get("display_name", ""),
            info.get("whatsapp", ""),
            info.get("telegram_id", ""),
            info.get("manager", ""),
            info.get("language", "ru"),
            "Да" if info.get("do_not_call") else "Нет",
            info.get("_debt_amount", ""),
            info.get("_days_overdue", ""),
            "Да" if info.get("_violation") else "",
            info.get("_registered_date", ""),
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            if col_idx == 1:
                cell.font = bold_font
            cell.alignment = Alignment(vertical="center")

    col_widths = [16, 45, 30, 16, 14, 14, 8, 12, 14, 14, 12, 16]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    try:
        wb.save(REGISTRY_XLSX)
        logger.info("Excel реестр сохранён: %s (%d записей)", REGISTRY_XLSX.name, len(contacts))
        return True
    except OSError as e:
        logger.error("Ошибка сохранения Excel реестра: %s", e)
        return False
