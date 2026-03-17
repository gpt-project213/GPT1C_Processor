#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tools/import_contacts.py
Импорт клиентской базы из Excel в config/debtors_contacts.json.

Использование:
  python tools/import_contacts.py                          # импорт из contacts.xlsx
  python tools/import_contacts.py --file contacts.xlsx     # явно указать файл
  python tools/import_contacts.py --file contacts.xlsx --overwrite  # перезаписать всё

Формат Excel (первая строка — заголовки):
  Клиент | Телефон | Email | Контактное лицо | Менеджер | Язык | Не звонить | Примечания

Логика:
  - Новые клиенты добавляются
  - Существующие обновляются (поля из Excel перезаписывают JSON)
  - --overwrite очищает JSON перед импортом

Версия: 1.0.0 (2026-03-17)
"""

import argparse
import json
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT))

try:
    import openpyxl
except ImportError:
    print("Ошибка: установите openpyxl — pip install openpyxl")
    sys.exit(1)

CONTACTS_PATH = _ROOT / "config" / "debtors_contacts.json"
DEFAULT_XLSX   = _ROOT / "contacts.xlsx"

# Колонки Excel → ключи JSON
# Поиск по нижнему регистру заголовка, допускаются синонимы
COLUMN_MAP = {
    "клиент":            "name",
    "контрагент":        "name",
    "наименование":      "name",
    "телефон":           "phone",
    "whatsapp":          "phone",       # если отдельной колонки нет — используем телефон
    "email":             "email",
    "e-mail":            "email",
    "контактное лицо":   "contact_person",
    "контакт":           "contact_person",
    "менеджер":          "manager",
    "язык":              "language",
    "не звонить":        "do_not_call",
    "do not call":       "do_not_call",
    "примечания":        "notes",
    "комментарий":       "notes",
}


def _parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    s = str(value).strip().lower()
    return s in ("да", "yes", "true", "1", "+")


def _load_existing() -> dict:
    if not CONTACTS_PATH.exists():
        return {}
    try:
        with open(CONTACTS_PATH, encoding="utf-8") as f:
            data = json.load(f)
        data.pop("_comment", None)
        return data
    except (OSError, json.JSONDecodeError) as e:
        print(f"Предупреждение: не удалось прочитать {CONTACTS_PATH}: {e}")
        return {}


def _save(data: dict) -> None:
    CONTACTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    out = {"_comment": "Справочник контактов должников. Ключ = имя клиента ТОЧНО как в debt JSON."}
    out.update(data)
    with open(CONTACTS_PATH, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)


def _empty_record(name: str) -> dict:
    return {
        "phone":               "",
        "email":               "",
        "contact_person":      "",
        "manager":             "",
        "language":            "ru",
        "do_not_call":         False,
        "notes":               "",
        "openclaw_session_id": None,
        "name_confirmations":  0,
        "phone_confirmations": 0,
    }


def import_xlsx(xlsx_path: Path, overwrite: bool = False) -> None:
    if not xlsx_path.exists():
        print(f"Файл не найден: {xlsx_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    # Читаем заголовки (первая строка)
    headers_raw = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
    # Маппинг: индекс колонки → ключ JSON
    col_to_key: dict[int, str] = {}
    whatsapp_col: int | None = None

    for i, h in enumerate(headers_raw):
        h_lower = h.lower()
        if h_lower in ("whatsapp", "вотсап", "ватсап"):
            whatsapp_col = i
        mapped = COLUMN_MAP.get(h_lower)
        if mapped:
            col_to_key[i] = mapped

    if not any(v == "name" for v in col_to_key.values()):
        print("Ошибка: колонка 'Клиент' / 'Контрагент' не найдена в Excel")
        sys.exit(1)

    existing = {} if overwrite else _load_existing()
    added = updated = skipped = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        # Пропускаем пустые строки
        if all(v is None or str(v).strip() == "" for v in row):
            continue

        # Извлекаем имя
        name = ""
        for i, key in col_to_key.items():
            if key == "name" and i < len(row) and row[i]:
                name = str(row[i]).strip()
                break
        if not name:
            skipped += 1
            continue

        is_new = name not in existing
        record = existing.get(name, _empty_record(name))
        # Убеждаемся что confirmations поля присутствуют (обратная совместимость)
        record.setdefault("name_confirmations", 0)
        record.setdefault("phone_confirmations", 0)

        for i, key in col_to_key.items():
            if key == "name" or i >= len(row):
                continue
            val = row[i]
            if val is None or str(val).strip() == "":
                continue
            if key == "do_not_call":
                record[key] = _parse_bool(val)
            else:
                record[key] = str(val).strip()

        # Отдельная колонка WhatsApp перезаписывает phone
        if whatsapp_col is not None and whatsapp_col < len(row):
            wa = row[whatsapp_col]
            if wa and str(wa).strip():
                record["phone"] = str(wa).strip()

        existing[name] = record
        if is_new:
            added += 1
        else:
            updated += 1

    _save(existing)
    total = len(existing)
    print(f"Готово: добавлено {added}, обновлено {updated}, пропущено {skipped}.")
    print(f"Итого в базе: {total} клиентов >> {CONTACTS_PATH}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Импорт контактов должников из Excel")
    parser.add_argument("--file", default=str(DEFAULT_XLSX),
                        help=f"Путь к Excel-файлу (по умолчанию: contacts.xlsx)")
    parser.add_argument("--overwrite", action="store_true",
                        help="Очистить базу перед импортом (по умолчанию: слияние)")
    args = parser.parse_args()

    import_xlsx(Path(args.file), overwrite=args.overwrite)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
