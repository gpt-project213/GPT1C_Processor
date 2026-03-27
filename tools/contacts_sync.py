#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tools/contacts_sync.py
Синхронизация базы клиентов между config/clients.json и Excel.

Версия: 1.0.0 (2026-03-27)

Использование:
  python tools/contacts_sync.py --export            # clients.json → contacts.xlsx
  python tools/contacts_sync.py --import            # contacts.xlsx → clients.json
  python tools/contacts_sync.py --export --file my.xlsx
  python tools/contacts_sync.py --import --file my.xlsx

Колонки Excel:
  Клиент (1С)       — ключ из 1С, не редактировать
  Менеджер          — привязка клиента к менеджеру
  Имя контакта      — display_name (как обращаться)
  WhatsApp          — номер для сообщений коллектора
  Адрес             — адрес торговой точки
  Язык              — ru / kz
  Не звонить        — Да / Нет
  Источники         — debt/sales (только чтение)
  Первое появление  — дата (только чтение)
  Последнее видели  — дата (только чтение)
"""

import argparse
import json
import logging
import os
import sys
import tempfile
from pathlib import Path

from dotenv import load_dotenv

ROOT_DIR = Path(__file__).resolve().parent.parent
load_dotenv(dotenv_path=ROOT_DIR / ".env", encoding="utf-8-sig", override=False)

sys.path.insert(0, str(ROOT_DIR))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Установите openpyxl: pip install openpyxl")
    sys.exit(1)

logging.basicConfig(format="%(asctime)s %(levelname)s %(message)s", level=logging.INFO)
logger = logging.getLogger("contacts_sync")

CONFIG_DIR = ROOT_DIR / "config"
CLIENTS_PATH = CONFIG_DIR / "clients.json"
DEFAULT_XLSX = ROOT_DIR / "contacts.xlsx"

# ─────────────────────────────────────────────────────────────
# Колонки
# ─────────────────────────────────────────────────────────────

COLUMNS = [
    ("Клиент (1С)",       "client_key",     False),  # (заголовок, поле, редактируемое)
    ("Менеджер",          "manager",        True),
    ("Имя контакта",      "display_name",   True),
    ("WhatsApp",          "whatsapp",       True),
    ("Адрес",             "address",        True),
    ("Язык",              "language",       True),
    ("Не звонить",        "do_not_call",    True),
    ("Источники",         "sources",        False),
    ("Первое появление",  "first_seen",     False),
    ("Последнее видели",  "last_seen",      False),
]

COL_HEADERS = [c[0] for c in COLUMNS]
COL_FIELDS  = [c[1] for c in COLUMNS]
COL_EDITABLE = [c[2] for c in COLUMNS]


# ─────────────────────────────────────────────────────────────
# Загрузка / сохранение clients.json
# ─────────────────────────────────────────────────────────────

def _load_clients() -> dict:
    if not CLIENTS_PATH.exists():
        return {"clients": {}}
    try:
        with open(CLIENTS_PATH, encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data.get("clients"), dict):
            data["clients"] = {}
        return data
    except (OSError, json.JSONDecodeError) as e:
        logger.error("Ошибка чтения clients.json: %s", e)
        return {"clients": {}}


def _save_clients(data: dict) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    try:
        tmp = tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8",
            dir=CONFIG_DIR, suffix=".tmp", delete=False,
        )
        json.dump(data, tmp, ensure_ascii=False, indent=2)
        tmp.close()
        os.replace(tmp.name, CLIENTS_PATH)
        logger.info("clients.json сохранён (%d клиентов)", len(data.get("clients", {})))
    except OSError as e:
        logger.error("Ошибка записи clients.json: %s", e)


# ─────────────────────────────────────────────────────────────
# Стили Excel
# ─────────────────────────────────────────────────────────────

def _header_style():
    return {
        "font": Font(bold=True, color="FFFFFF", size=11),
        "fill": PatternFill("solid", fgColor="1A3A5C"),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }

def _readonly_fill():
    return PatternFill("solid", fgColor="F0F0F0")

def _editable_fill():
    return PatternFill("solid", fgColor="FFFFFF")

def _border():
    side = Side(style="thin", color="CCCCCC")
    return Border(left=side, right=side, top=side, bottom=side)


# ─────────────────────────────────────────────────────────────
# ЭКСПОРТ: clients.json → Excel
# ─────────────────────────────────────────────────────────────

def export_to_excel(xlsx_path: Path) -> int:
    """Экспортирует clients.json в Excel. Возвращает количество строк."""
    data = _load_clients()
    clients_db = data.get("clients", {})

    wb = openpyxl.Workbook()

    # ── Лист 1: Контакты ─────────────────────────────────────
    ws = wb.active
    ws.title = "Контакты клиентов"

    # Заголовки
    for col_idx, header in enumerate(COL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        for attr, val in _header_style().items():
            setattr(cell, attr, val)
        cell.border = _border()

    ws.row_dimensions[1].height = 30

    # Данные — сортировка по менеджеру, потом по имени клиента
    sorted_clients = sorted(
        clients_db.items(),
        key=lambda x: (x[1].get("manager", ""), x[0])
    )

    for row_idx, (client_key, info) in enumerate(sorted_clients, start=2):
        if not isinstance(info, dict):
            continue

        sources = ", ".join(info.get("sources", []))
        do_not_call = "Да" if info.get("do_not_call") else "Нет"

        row_values = [
            client_key,
            info.get("manager", ""),
            info.get("display_name", ""),
            info.get("whatsapp", ""),
            info.get("address", ""),
            info.get("language", "ru"),
            do_not_call,
            sources,
            info.get("first_seen", ""),
            info.get("last_seen", ""),
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = _border()
            if COL_EDITABLE[col_idx - 1]:
                cell.fill = _editable_fill()
            else:
                cell.fill = _readonly_fill()
                cell.font = Font(color="888888")

    # Ширина колонок
    col_widths = [40, 12, 20, 18, 30, 8, 12, 14, 16, 16]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    # ── Лист 2: Инструкция ────────────────────────────────────
    ws2 = wb.create_sheet("Инструкция")
    instructions = [
        ("Поле", "Описание", "Пример"),
        ("Клиент (1С)", "Имя ТОЧНО как в 1С. НЕ редактировать — ключ для синхронизации", "ТОО Альфа Трейд"),
        ("Менеджер", "Менеджер, ведущий клиента", "Алена"),
        ("Имя контакта", "Как обращаться к контактному лицу клиента", "Аида"),
        ("WhatsApp", "Номер для WhatsApp-сообщений (87xxxxxxxxx)", "87012345678"),
        ("Адрес", "Адрес торговой точки", "ул. Достык 12, магазин Аида"),
        ("Язык", "Язык сообщений: ru или kz", "ru"),
        ("Не звонить", "Да — только WhatsApp/Telegram, звонки запрещены", "Нет"),
        ("Источники", "Откуда появился клиент: debt/sales (НЕ редактировать)", "debt, sales"),
        ("Первое появление", "Дата первого появления в системе (НЕ редактировать)", "2026-03-25"),
        ("Последнее видели", "Дата последнего обновления (НЕ редактировать)", "2026-03-27"),
        ("", "", ""),
        ("⚠️ Важно", "Серые колонки (Клиент 1С, Источники, Даты) — не редактировать", ""),
        ("⚠️ Важно", "После заполнения запустить: python tools/contacts_sync.py --import", ""),
        ("⚠️ Важно", "Принадлежность клиента менеджеру — колонка Менеджер", ""),
    ]
    for row_data in instructions:
        ws2.append(row_data)

    for cell in ws2[1]:
        cell.font = Font(bold=True)

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 60
    ws2.column_dimensions["C"].width = 25

    wb.save(xlsx_path)
    count = len(sorted_clients)
    logger.info("Экспорт завершён: %d клиентов → %s", count, xlsx_path)
    return count


# ─────────────────────────────────────────────────────────────
# ИМПОРТ: Excel → clients.json
# ─────────────────────────────────────────────────────────────

def import_from_excel(xlsx_path: Path) -> dict:
    """
    Импортирует данные из Excel в clients.json.
    Обновляет только редактируемые поля (manager, display_name, whatsapp,
    address, language, do_not_call).
    Новых клиентов НЕ создаёт — только обновляет существующих.

    Возвращает статистику: {"updated": N, "skipped": N, "not_found": N}
    """
    if not xlsx_path.exists():
        logger.error("Файл не найден: %s", xlsx_path)
        return {"updated": 0, "skipped": 0, "not_found": 0}

    try:
        wb = openpyxl.load_workbook(xlsx_path)
    except Exception as e:
        logger.error("Ошибка открытия Excel: %s", e)
        return {"updated": 0, "skipped": 0, "not_found": 0}

    # Ищем лист с данными
    sheet_name = None
    for name in wb.sheetnames:
        if "контакт" in name.lower() or name == wb.sheetnames[0]:
            sheet_name = name
            break
    ws = wb[sheet_name]

    # Читаем заголовки из первой строки
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    try:
        idx_key     = headers.index("Клиент (1С)")
        idx_manager = headers.index("Менеджер")
        idx_name    = headers.index("Имя контакта")
        idx_wa      = headers.index("WhatsApp")
        idx_addr    = headers.index("Адрес") if "Адрес" in headers else -1
        idx_lang    = headers.index("Язык")
        idx_dnc     = headers.index("Не звонить")
    except ValueError as e:
        logger.error("Не найдена колонка в Excel: %s", e)
        return {"updated": 0, "skipped": 0, "not_found": 0}

    data = _load_clients()
    clients_db = data.get("clients", {})

    stats = {"updated": 0, "skipped": 0, "not_found": 0}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[idx_key] is None:
            continue

        client_key = str(row[idx_key]).strip()
        if not client_key:
            continue

        if client_key not in clients_db:
            logger.warning("Клиент не найден в базе: %s", client_key)
            stats["not_found"] += 1
            continue

        entry = clients_db[client_key]
        changed = False

        def _str(val):
            return str(val).strip() if val is not None else ""

        # Менеджер
        manager = _str(row[idx_manager])
        if manager and entry.get("manager") != manager:
            entry["manager"] = manager
            changed = True

        # Имя контакта (display_name)
        display = _str(row[idx_name])
        if display and entry.get("display_name") != display:
            entry["display_name"] = display
            changed = True

        # WhatsApp
        wa = _str(row[idx_wa])
        # Нормализация: убираем лишние символы, оставляем цифры и +
        wa_clean = "".join(c for c in wa if c.isdigit() or c == "+")
        if wa_clean and entry.get("whatsapp") != wa_clean:
            entry["whatsapp"] = wa_clean
            changed = True

        # Адрес
        if idx_addr >= 0:
            addr = _str(row[idx_addr])
            if addr and entry.get("address") != addr:
                entry["address"] = addr
                changed = True

        # Язык
        lang = _str(row[idx_lang]).lower()
        if lang in ("ru", "kz") and entry.get("language") != lang:
            entry["language"] = lang
            changed = True

        # Не звонить
        dnc_val = _str(row[idx_dnc]).lower()
        dnc = dnc_val in ("да", "yes", "1", "true")
        if entry.get("do_not_call") != dnc:
            entry["do_not_call"] = dnc
            changed = True

        if changed:
            stats["updated"] += 1
            logger.info("Обновлён: %s (менеджер=%s, wa=%s)", client_key, manager or "—", wa_clean or "—")
        else:
            stats["skipped"] += 1

    data["clients"] = clients_db
    _save_clients(data)

    logger.info(
        "Импорт завершён: обновлено %d, без изменений %d, не найдено %d",
        stats["updated"], stats["skipped"], stats["not_found"]
    )
    return stats


# ─────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Синхронизация базы клиентов: clients.json ↔ Excel"
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--export", action="store_true", help="clients.json → Excel")
    group.add_argument("--import", dest="do_import", action="store_true",
                       help="Excel → clients.json")
    parser.add_argument("--file", type=Path, default=DEFAULT_XLSX,
                        help=f"Путь к Excel-файлу (по умолчанию: {DEFAULT_XLSX})")
    args = parser.parse_args()

    if args.export:
        count = export_to_excel(args.file)
        print(f"✅ Экспорт: {count} клиентов → {args.file}")
    else:
        stats = import_from_excel(args.file)
        print(
            f"✅ Импорт: обновлено {stats['updated']}, "
            f"без изменений {stats['skipped']}, "
            f"не найдено {stats['not_found']}"
        )


if __name__ == "__main__":
    main()
